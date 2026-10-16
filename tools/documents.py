"""Document tools: create PowerPoint (.pptx) and Word (.docx) files in workspace/output.

Files are always NEW (never overwritten) and never leave the workspace sandbox.
Bangla text automatically gets a Bangla-capable font so it renders correctly.
Presentations are 16:9 with a designed theme ("academic" navy or "modern" teal).
"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from tools.files import safe_path
from tools.registry import Permission, ToolError, ToolSpec

BANGLA_FONT = "Nirmala UI"
_BANGLA = re.compile(r"[\u0980-\u09FF]")
MAX_SLIDES = 40
MAX_BULLETS = 8
MAX_SECTIONS = 60
MAX_ITEMS = 40
MAX_TEXT = 3000

SLIDE_W, SLIDE_H = 13.333, 7.5
THEMES: dict[str, dict[str, str]] = {
    # Navy + gold: clean look for university work.
    "academic": {"dark": "1F3A5F", "deco": "2B4A75", "accent": "C9A227", "bg": "FFFFFF", "text": "1B2733",
                 "muted": "6B7785", "sub": "C9D6E8", "card_fill": "2E6FBA", "card_text": "FFFFFF"},
    # Teal + green: colourful and modern.
    "modern": {"dark": "0B4F4A", "deco": "11695F", "accent": "34D399", "bg": "F2FBF9", "text": "10302D",
               "muted": "5B7C78", "sub": "B7EDE4", "card_fill": "34D399", "card_text": "062F2B"},
}


def _has_bangla(text: str) -> bool:
    return bool(_BANGLA.search(text or ""))


def _text(value: Any, limit: int = MAX_TEXT) -> str:
    return ("" if value is None else str(value)).strip()[:limit]


def _text_list(values: Any, max_items: int) -> list[str]:
    if values is None:
        return []
    if isinstance(values, str):
        values = [values]
    try:
        items = list(values)[:max_items]
    except TypeError:
        return []
    return [t for t in (_text(v) for v in items) if t]


def _as_dict(item: Any, key: str) -> dict[str, Any]:
    """Models sometimes send a bare string instead of an object; accept both."""
    return item if isinstance(item, dict) else {key: item}


def _new_path(output: Path, filename: str, ext: str) -> Path:
    name = _text(filename, 120)
    if not name:
        raise ToolError("filename is required")
    if not name.lower().endswith(ext):
        if Path(name).suffix:
            raise ToolError(f"File name must end with {ext}")
        name += ext
    target = safe_path(output, name)
    if target.exists():
        raise ToolError("File already exists. Choose a new file name (existing files are never overwritten).")
    target.parent.mkdir(parents=True, exist_ok=True)
    return target


# ----------------------------------------------------------------- PowerPoint
def _rgb(hex_color: str) -> Any:
    from pptx.dml.color import RGBColor

    return RGBColor.from_string(hex_color)


def _format_frame(frame: Any, size: float | None = None, color: str | None = None,
                  bold: bool | None = None, align: Any = None, space_after: float | None = None) -> None:
    """Apply size/colour/weight to every run; give Bangla runs a Bangla-capable font."""
    from pptx.oxml.ns import qn
    from pptx.util import Pt

    for paragraph in frame.paragraphs:
        if align is not None:
            paragraph.alignment = align
        if space_after is not None:
            paragraph.space_after = Pt(space_after)
        for run in paragraph.runs:
            if size:
                run.font.size = Pt(size)
            if bold is not None:
                run.font.bold = bold
            if color:
                run.font.color.rgb = _rgb(color)
            if _has_bangla(run.text):
                run.font.name = BANGLA_FONT
                props = run._r.get_or_add_rPr()
                for old in props.findall(qn("a:cs")):
                    props.remove(old)
                props.append(props.makeelement(qn("a:cs"), {"typeface": BANGLA_FONT}))


def _place(shape: Any, x: float, y: float, w: float, h: float) -> None:
    from pptx.util import Inches

    shape.left, shape.top, shape.width, shape.height = Inches(x), Inches(y), Inches(w), Inches(h)


def _background(slide: Any, color: str) -> None:
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _rgb(color)


def _decor(slide: Any, shapes: list[tuple[Any, float, float, float, float, str]]) -> None:
    """Add decorative shapes BEHIND the placeholders, bottom-to-top in list order."""
    from pptx.util import Inches

    tree = slide.shapes._spTree
    for index, (kind, x, y, w, h, fill) in enumerate(shapes):
        shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(fill)
        shape.line.fill.background()
        shape.shadow.inherit = False
        tree.remove(shape._element)
        tree.insert(2 + index, shape._element)


def _textbox(slide: Any, text: str, x: float, y: float, w: float, h: float, size: float,
             color: str, align: Any = None, bold: bool = False) -> None:
    from pptx.util import Inches

    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.text_frame.text = text
    box.text_frame.word_wrap = True
    _format_frame(box.text_frame, size, color, bold, align)


def _fit_size(bullets: list[str], width_in: float, height_in: float) -> int:
    """Largest font size (pt) at which the bullets should fit the body box."""
    usable_w = max(2.0, width_in - 0.6)  # bullet indent
    for size in (28, 26, 24, 22, 20, 18, 16, 14):
        chars_per_line = max(10, int(usable_w * 72 / (size * 0.38)))
        lines = sum(max(1, -(-len(b) // chars_per_line)) for b in bullets)
        needed = lines * size * 1.45 + len(bullets) * 10
        if needed <= height_in * 72 * 0.95:
            return size
    return 14


def _size_for_title(text: str, big: int, mid: int, small: int, limits: tuple[int, int]) -> int:
    return big if len(text) <= limits[0] else mid if len(text) <= limits[1] else small


def _cover_slide(prs: Any, th: dict[str, str], modern: bool, title: str, subtitle: str) -> None:
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

    slide = prs.slides.add_slide(prs.slide_layouts[0])
    _background(slide, th["dark"])
    if modern:
        _decor(slide, [(MSO_SHAPE.OVAL, 9.0, -1.8, 7.0, 7.0, th["deco"]),
                       (MSO_SHAPE.OVAL, 10.6, 4.3, 4.0, 4.0, "0F766E"),
                       (MSO_SHAPE.RECTANGLE, 0.95, 4.3, 1.8, 0.1, th["accent"])])
        width = 8.0
    else:
        _decor(slide, [(MSO_SHAPE.RECTANGLE, 0, 0, 0.4, SLIDE_H, th["accent"]),
                       (MSO_SHAPE.RECTANGLE, 0.95, 4.3, 1.8, 0.1, th["accent"])])
        width = 11.4
    slide.shapes.title.text = title
    _place(slide.shapes.title, 0.95, 1.4, width, 2.7)
    slide.shapes.title.text_frame.vertical_anchor = MSO_ANCHOR.BOTTOM
    slide.shapes.title.text_frame.word_wrap = True
    _format_frame(slide.shapes.title.text_frame, _size_for_title(title, 48, 40, 32, (28, 50)), "FFFFFF", True, PP_ALIGN.LEFT)
    subtitle_ph = slide.placeholders[1]
    if subtitle:
        subtitle_ph.text = subtitle
        _place(subtitle_ph, 0.95, 4.6, width, 1.4)
        subtitle_ph.text_frame.vertical_anchor = MSO_ANCHOR.TOP
        _format_frame(subtitle_ph.text_frame, 20, th["sub"], False, PP_ALIGN.LEFT)
    else:
        subtitle_ph._element.getparent().remove(subtitle_ph._element)


def _content_slide(prs: Any, th: dict[str, str], deck_title: str, number: int, item: dict[str, Any]) -> None:
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches

    slide_title = _text(item.get("title"), 200) or "Untitled"
    bullets = _text_list(item.get("bullets"), MAX_BULLETS)
    highlight = _text(item.get("highlight"), 140)
    notes = _text(item.get("notes"))

    slide = prs.slides.add_slide(prs.slide_layouts[1])
    _background(slide, th["bg"])
    decor = [(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, 1.35, th["dark"]),
             (MSO_SHAPE.RECTANGLE, 0, 1.35, SLIDE_W, 0.07, th["accent"])]
    if bullets:
        decor.append((MSO_SHAPE.RECTANGLE, 0.72, 1.95, 0.07, 4.6, th["accent"]))
    _decor(slide, decor)

    slide.shapes.title.text = slide_title
    _place(slide.shapes.title, 0.8, 0.15, 11.7, 1.05)
    slide.shapes.title.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    slide.shapes.title.text_frame.word_wrap = True
    _format_frame(slide.shapes.title.text_frame, _size_for_title(slide_title, 32, 28, 24, (40, 65)), "FFFFFF", True, PP_ALIGN.LEFT)

    body = slide.placeholders[1]
    if bullets:
        width = 7.9 if highlight else 11.5
        _place(body, 1.0, 1.85, width, 4.85)
        frame = body.text_frame
        frame.word_wrap = True
        for i, bullet in enumerate(bullets):
            paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            paragraph.text = bullet
        _format_frame(frame, _fit_size(bullets, width, 4.85), th["text"], False, PP_ALIGN.LEFT, space_after=10)
    else:
        body._element.getparent().remove(body._element)

    if highlight:
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.3), Inches(2.0), Inches(3.5), Inches(3.9))
        card.adjustments[0] = 0.07
        card.fill.solid()
        card.fill.fore_color.rgb = _rgb(th["card_fill"])
        card.line.fill.background()
        card.shadow.inherit = False
        frame = card.text_frame
        frame.text = highlight
        frame.word_wrap = True
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        size = 40 if len(highlight) <= 18 else 32 if len(highlight) <= 40 else 24
        _format_frame(frame, size, th["card_text"], True, PP_ALIGN.CENTER)

    _textbox(slide, deck_title, 0.8, 6.95, 9.5, 0.35, 11, th["muted"], PP_ALIGN.LEFT)
    _textbox(slide, str(number), 11.5, 6.95, 1.1, 0.35, 11, th["muted"], PP_ALIGN.RIGHT)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes


def _closing_slide(prs: Any, th: dict[str, str], modern: bool, text: str) -> None:
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    _background(slide, th["dark"])
    shapes = [(MSO_SHAPE.RECTANGLE, 5.8, 4.55, 1.7, 0.1, th["accent"])]
    if modern:
        shapes = [(MSO_SHAPE.OVAL, -2.0, 4.2, 6.0, 6.0, th["deco"]), (MSO_SHAPE.OVAL, 10.5, -2.2, 5.5, 5.5, th["deco"])] + shapes
    _decor(slide, shapes)
    slide.shapes.title.text = text
    _place(slide.shapes.title, 1.0, 2.4, 11.3, 2.0)
    slide.shapes.title.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    slide.shapes.title.text_frame.word_wrap = True
    _format_frame(slide.shapes.title.text_frame, _size_for_title(text, 56, 44, 34, (24, 45)), "FFFFFF", True, PP_ALIGN.CENTER)


# ----------------------------------------------------------------------- Word
def _style_docx(doc: Any) -> None:
    from docx.oxml.ns import qn

    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if _has_bangla(run.text):
                run.font.name = BANGLA_FONT
                fonts = run._element.get_or_add_rPr().rFonts
                if fonts is not None:
                    fonts.set(qn("w:cs"), BANGLA_FONT)


def make_document_tools(workspace: Path) -> list[ToolSpec]:
    """Build create_pptx / create_docx bound to one workspace directory."""
    workspace = workspace.resolve()
    output = workspace / "output"

    def create_pptx(filename: str, title: str, slides: Any, subtitle: str = "",
                    theme: str = "academic", closing: str = "") -> dict[str, Any]:
        try:
            from pptx import Presentation
            from pptx.util import Inches
        except ImportError as exc:
            raise ToolError("python-pptx is not installed on the server") from exc
        items = list(slides) if isinstance(slides, (list, tuple)) else []
        if not items:
            raise ToolError("Provide at least one slide in 'slides'")
        if len(items) > MAX_SLIDES:
            raise ToolError(f"Too many slides (max {MAX_SLIDES})")
        theme_name = (_text(theme, 20) or "academic").lower()
        if theme_name not in THEMES:
            raise ToolError(f"Unknown theme '{theme_name}'. Use one of: {sorted(THEMES)}")
        th = THEMES[theme_name]
        modern = theme_name == "modern"
        deck_title = _text(title, 200) or "Presentation"
        target = _new_path(output, filename, ".pptx")

        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)
        _cover_slide(prs, th, modern, deck_title, _text(subtitle, 300))
        for number, raw in enumerate(items, start=2):
            _content_slide(prs, th, deck_title, number, _as_dict(raw, "title"))
        closing_text = _text(closing, 80)
        if closing_text:
            _closing_slide(prs, th, modern, closing_text)
        prs.core_properties.title = deck_title
        prs.save(str(target))
        return {"saved": str(target.relative_to(workspace)), "slides": len(prs.slides), "theme": theme_name}

    def create_docx(filename: str, title: str, sections: Any, author: str = "") -> dict[str, Any]:
        try:
            from docx import Document
        except ImportError as exc:
            raise ToolError("python-docx is not installed on the server") from exc
        items = list(sections) if isinstance(sections, (list, tuple)) else []
        if not items:
            raise ToolError("Provide at least one section in 'sections'")
        if len(items) > MAX_SECTIONS:
            raise ToolError(f"Too many sections (max {MAX_SECTIONS})")
        doc_title = _text(title, 200) or "Document"
        target = _new_path(output, filename, ".docx")

        doc = Document()
        doc.core_properties.title = doc_title
        if _text(author, 100):
            doc.core_properties.author = _text(author, 100)
        doc.add_heading(doc_title, level=0)
        for raw in items:
            item = _as_dict(raw, "heading")
            heading = _text(item.get("heading"), 200)
            if heading:
                doc.add_heading(heading, level=1)
            for para in _text_list(item.get("paragraphs"), MAX_ITEMS):
                doc.add_paragraph(para)
            for bullet in _text_list(item.get("bullets"), MAX_ITEMS):
                doc.add_paragraph(bullet, style="List Bullet")
        _style_docx(doc)
        doc.save(str(target))
        return {"saved": str(target.relative_to(workspace)), "sections": len(items)}

    return [
        ToolSpec(
            name="create_pptx",
            description=(
                "Create a NEW designed 16:9 PowerPoint (.pptx) in workspace/output. Never overwrites existing files. "
                "Build the WHOLE deck in one call with rich, specific content: for a university talk use 8-12 content slides "
                "(introduction/definition, background, causes, effects, case study or examples, solutions, conclusion). "
                "Each content slide needs 4-6 bullets written as full informative sentences (12-25 words) with concrete facts or examples, "
                "plus speaker notes (2-4 sentences). Put a short key fact or number in 'highlight' on at least 3 slides. "
                "Only use statistics you are sure about or found with web_search; never invent numbers or sources. "
                "Use theme 'academic' for university/school work and 'modern' for creative or informal decks."
            ),
            parameters={
                "type": "object",
                "properties": {
                    "filename": {"type": "string", "description": "File name, e.g. 'river_erosion.pptx'"},
                    "title": {"type": "string", "description": "Title shown on the cover slide"},
                    "subtitle": {"type": "string", "description": "Cover subtitle (course, presenter, date)"},
                    "theme": {"type": "string", "description": "'academic' (navy and gold) or 'modern' (teal and green)", "enum": ["academic", "modern"]},
                    "closing": {"type": "string", "description": "Text for the final slide, e.g. 'ধন্যবাদ' or 'Thank you'"},
                    "slides": {
                        "type": "array",
                        "description": "Content slides in order (max 40)",
                        "items": {
                            "type": "object",
                            "properties": {
                                "title": {"type": "string", "description": "Slide title"},
                                "bullets": {"type": "array", "items": {"type": "string"}, "description": "4-6 full-sentence bullet points"},
                                "highlight": {"type": "string", "description": "Optional key fact or number shown in a coloured card (max ~12 words)"},
                                "notes": {"type": "string", "description": "Speaker notes, 2-4 sentences"},
                            },
                            "required": ["title"],
                        },
                    },
                },
                "required": ["filename", "title", "slides"],
            },
            func=create_pptx, permission=Permission.SAFE_WRITE, timeout=30,
        ),
        ToolSpec(
            name="create_docx",
            description=(
                "Create a NEW Word (.docx) document in workspace/output. "
                "Write real, complete text. Never overwrites existing files."
            ),
            parameters={
                "type": "object",
                "properties": {
                    "filename": {"type": "string", "description": "File name, e.g. 'assignment.docx'"},
                    "title": {"type": "string", "description": "Document title"},
                    "author": {"type": "string", "description": "Optional author name"},
                    "sections": {
                        "type": "array",
                        "description": "Sections in order (max 60)",
                        "items": {
                            "type": "object",
                            "properties": {
                                "heading": {"type": "string", "description": "Section heading"},
                                "paragraphs": {"type": "array", "items": {"type": "string"}, "description": "Body paragraphs"},
                                "bullets": {"type": "array", "items": {"type": "string"}, "description": "Bullet points"},
                            },
                            "required": ["heading"],
                        },
                    },
                },
                "required": ["filename", "title", "sections"],
            },
            func=create_docx, permission=Permission.SAFE_WRITE, timeout=30,
        ),
    ]

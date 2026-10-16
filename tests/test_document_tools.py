import pytest
from docx import Document
from pptx import Presentation
from pptx.util import Inches

from tools import build_registry
from tools.documents import BANGLA_FONT, _fit_size


@pytest.fixture
def reg(tmp_path):
    (tmp_path / "input").mkdir()
    (tmp_path / "output").mkdir()
    return build_registry(tmp_path)


def _texts(slide):
    return [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]


def test_pptx_roundtrip_with_bangla(reg, tmp_path):
    res = reg.execute("create_pptx", {
        "filename": "demo", "title": "Demo deck", "subtitle": "Test",
        "slides": [
            {"title": "One", "bullets": ["a", "b"], "notes": "speaker note"},
            {"title": "নদী ভাঙন", "bullets": ["কারণ", "ফলাফল"]},
        ],
    })
    assert res["ok"], res
    assert res["result"]["saved"] == "output/demo.pptx"
    prs = Presentation(str(tmp_path / "output" / "demo.pptx"))
    assert len(prs.slides) == 3
    assert prs.slides[1].notes_slide.notes_text_frame.text == "speaker note"
    runs = [r for sh in prs.slides[2].shapes if sh.has_text_frame for p in sh.text_frame.paragraphs for r in p.runs]
    assert any(r.font.name == BANGLA_FONT for r in runs)


@pytest.mark.parametrize("theme", ["academic", "modern"])
def test_pptx_design_features(reg, tmp_path, theme):
    res = reg.execute("create_pptx", {
        "filename": f"deck_{theme}", "title": "Cover", "subtitle": "Sub", "theme": theme, "closing": "ধন্যবাদ",
        "slides": [{"title": "S1", "bullets": ["x"], "highlight": "KEY FACT 70%"}, {"title": "S2", "bullets": ["y"]}],
    })
    assert res["ok"], res
    assert res["result"]["slides"] == 4 and res["result"]["theme"] == theme
    prs = Presentation(str(tmp_path / "output" / f"deck_{theme}.pptx"))
    assert prs.slide_width == Inches(13.333) and prs.slide_height == Inches(7.5)
    assert "KEY FACT 70%" in _texts(prs.slides[1])
    assert "ধন্যবাদ" in _texts(prs.slides[3])
    assert "2" in _texts(prs.slides[1]) and "3" in _texts(prs.slides[2])  # slide numbers


def test_unknown_theme_rejected(reg):
    res = reg.execute("create_pptx", {"filename": "t", "title": "T", "theme": "neon", "slides": [{"title": "S"}]})
    assert not res["ok"] and "Unknown theme" in res["error"]


def test_fit_size_shrinks_for_more_text():
    short = _fit_size(["a short line"] * 3, 11.5, 4.85)
    long = _fit_size(["x" * 150] * 8, 7.9, 4.85)
    assert short > long >= 14


def test_docx_roundtrip(reg, tmp_path):
    res = reg.execute("create_docx", {
        "filename": "essay.docx", "title": "My Essay", "author": "Mee",
        "sections": [{"heading": "Intro", "paragraphs": ["Hello world"], "bullets": ["one", "two"]},
                     {"heading": "সারসংক্ষেপ", "paragraphs": ["বাংলা লেখা"]}],
    })
    assert res["ok"], res
    doc = Document(str(tmp_path / "output" / "essay.docx"))
    texts = [p.text for p in doc.paragraphs]
    assert "My Essay" in texts and "Hello world" in texts and "বাংলা লেখা" in texts
    assert sum(1 for p in doc.paragraphs if p.style.name == "List Bullet") == 2


def test_never_overwrites(reg):
    args = {"filename": "x.docx", "title": "T", "sections": [{"heading": "H"}]}
    assert reg.execute("create_docx", args)["ok"]
    again = reg.execute("create_docx", args)
    assert not again["ok"] and "already exists" in again["error"]


@pytest.mark.parametrize("name", ["../escape.pptx", "/etc/x.pptx", "notes.txt", ""])
def test_bad_file_names_rejected(reg, name):
    res = reg.execute("create_pptx", {"filename": name, "title": "T", "slides": [{"title": "S"}]})
    assert not res["ok"]


def test_empty_and_oversized_input_rejected(reg):
    assert not reg.execute("create_pptx", {"filename": "a", "title": "T", "slides": []})["ok"]
    assert not reg.execute("create_pptx", {"filename": "b", "title": "T", "slides": [{"title": "s"}] * 41})["ok"]
    assert not reg.execute("create_docx", {"filename": "c", "title": "T", "sections": []})["ok"]


def test_lenient_with_sloppy_model_output(reg):
    res = reg.execute("create_pptx", {"filename": "lenient", "title": "T", "slides": ["just a title", {"title": "B", "bullets": "single string"}]})
    assert res["ok"], res

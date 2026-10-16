import pytest
from openpyxl import load_workbook

from agent.permissions import describe_action
from tools import build_registry

SHEETS = [{"name": "Budget", "columns": ["Item", "Cost", "Code"],
           "rows": [["চা", "1250", "007"], ["Rent", "3.5", "=1+1"], ["Total", "1,250", "@x"]]}]


@pytest.fixture
def reg(tmp_path):
    (tmp_path / "output").mkdir()
    return build_registry(tmp_path)


def test_workbook_roundtrip_and_safe_values(reg, tmp_path):
    res = reg.execute("create_xlsx", {"filename": "budget", "sheets": SHEETS})
    assert res["ok"], res
    assert res["result"] == {"saved": "output/budget.xlsx", "sheets": 1, "rows": 3}
    sheet = load_workbook(tmp_path / "output" / "budget.xlsx")["Budget"]
    assert [c.value for c in sheet[1]] == ["Item", "Cost", "Code"]
    assert sheet["A2"].value == "চা" and sheet["B2"].value == 1250 and isinstance(sheet["B2"].value, int)
    assert sheet["C2"].value == "007"                     # identifiers stay text
    assert sheet["B3"].value == 3.5
    assert sheet["C3"].value == "=1+1" and sheet["C3"].data_type == "s"   # never a formula
    assert sheet["B4"].value == "1,250"                   # thousands separators stay text
    assert sheet["A1"].font.bold and sheet.freeze_panes == "A2" and sheet.auto_filter.ref == "A1:C4"


def test_sheet_names_are_cleaned_and_unique(reg, tmp_path):
    sheets = [{"name": "A/B:C?", "columns": ["x"], "rows": [["1"]]},
              {"name": "A/B:C?", "columns": ["x"], "rows": [["2"]]},
              {"name": "N" * 50, "columns": ["x"], "rows": []}]
    assert reg.execute("create_xlsx", {"filename": "names", "sheets": sheets, "theme": "modern"})["ok"]
    names = load_workbook(tmp_path / "output" / "names.xlsx").sheetnames
    assert len(names) == 3 and len(set(n.lower() for n in names)) == 3 and all(len(n) <= 31 for n in names)
    assert not any(ch in n for n in names for ch in "[]:*?/\\")


def test_limits_and_bad_input_rejected(reg):
    ok_sheet = {"name": "S", "columns": ["a"], "rows": [["1"]]}
    assert not reg.execute("create_xlsx", {"filename": "a", "sheets": []})["ok"]
    assert not reg.execute("create_xlsx", {"filename": "b", "sheets": [ok_sheet] * 11})["ok"]
    assert not reg.execute("create_xlsx", {"filename": "c", "sheets": [{"name": "S", "columns": ["a"], "rows": [["1"]] * 5001}]})["ok"]
    assert not reg.execute("create_xlsx", {"filename": "d", "sheets": [{"name": "S", "columns": [], "rows": []}]})["ok"]
    assert not reg.execute("create_xlsx", {"filename": "e", "sheets": [{"name": "S", "columns": ["c"] * 51, "rows": []}]})["ok"]
    assert not reg.execute("create_xlsx", {"filename": "f", "sheets": [ok_sheet], "theme": "neon"})["ok"]
    for bad in ["../x.xlsx", "/etc/x.xlsx", "x.csv", ""]:
        assert not reg.execute("create_xlsx", {"filename": bad, "sheets": [ok_sheet]})["ok"]


def test_never_overwrites(reg):
    args = {"filename": "once", "sheets": [{"name": "S", "columns": ["a"], "rows": []}]}
    assert reg.execute("create_xlsx", args)["ok"]
    again = reg.execute("create_xlsx", args)
    assert not again["ok"] and "already exists" in again["error"]


def test_lenient_with_sloppy_model_output(reg, tmp_path):
    sheets = [{"name": "S", "columns": ["a", "b"], "rows": [["1", "x"], "lonely", ["2"]]}]
    assert reg.execute("create_xlsx", {"filename": "loose", "sheets": sheets})["ok"]


def test_approval_card_text():
    title, details = describe_action("create_xlsx", {"filename": "budget.xlsx", "sheets": SHEETS})
    assert "budget.xlsx" in title and details[0] == "1 sheets" and "- Budget: 3 rows" in details


def test_create_xlsx_requires_approval(reg):
    from agent.permissions import apply_approval_policy

    assert apply_approval_policy(reg).get("create_xlsx").permission.name == "CONFIRM_REQUIRED"
